"""Renders a generated question set into an official-looking .docx paper."""
import os

from docx import Document
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.shared import Cm, Pt

SECTION_TITLES = {
    "descriptive": "Descriptive Questions",
    "fill_blank": "Fill in the Blanks",
    "mcq": "Multiple Choice Questions",
}
OPTION_LETTERS = ["A", "B", "C", "D", "E", "F"]

PAGE_MARGIN = Cm(2)
USABLE_WIDTH = Cm(21) - PAGE_MARGIN - PAGE_MARGIN  # A4 width minus left/right margins
MARK_COL_WIDTH = Cm(2.2)
QUESTION_COL_WIDTH = USABLE_WIDTH - MARK_COL_WIDTH


def _set_a4_page(doc):
    section = doc.sections[0]
    section.page_width = Cm(21)
    section.page_height = Cm(29.7)
    section.left_margin = PAGE_MARGIN
    section.right_margin = PAGE_MARGIN
    section.top_margin = Cm(1.8)
    section.bottom_margin = Cm(1.8)


def _set_font(run, size=11, bold=False):
    run.font.size = Pt(size)
    run.font.bold = bold


def _header(doc, meta, course_name, fmt, date_str, set_label):
    p = doc.add_paragraph()
    p.alignment = WD_ALIGN_PARAGRAPH.CENTER
    r = p.add_run(meta.get("institute", "Diesel Traction Training Centre, Kazipet (DTTC/KZJ)"))
    _set_font(r, 15, True)

    p2 = doc.add_paragraph()
    p2.alignment = WD_ALIGN_PARAGRAPH.CENTER
    r2 = p2.add_run(meta.get("zone", "South Central Railway"))
    _set_font(r2, 12, True)

    p3 = doc.add_paragraph()
    p3.alignment = WD_ALIGN_PARAGRAPH.CENTER
    r3 = p3.add_run(f"Question Paper — {fmt['name']}")
    _set_font(r3, 12, True)

    doc.add_paragraph()

    total_marks = sum(
        (sec["display_count"] - sec.get("choice_leave", 0)) * sec["marks_each"]
        for sec in fmt["sections"]
    )
    info = doc.add_table(rows=2, cols=2)
    info.autofit = True
    cells = [
        (f"Course: {course_name}", f"Set: {set_label}"),
        (f"Date: {date_str}", f"Maximum Marks: {total_marks}"),
    ]
    for r_idx, (left, right) in enumerate(cells):
        lc = info.cell(r_idx, 0).paragraphs[0].add_run(left)
        _set_font(lc, 11, False)
        rc = info.cell(r_idx, 1).paragraphs[0].add_run(right)
        _set_font(rc, 11, False)

    doc.add_paragraph()
    line = doc.add_paragraph("_" * 90)
    line.alignment = WD_ALIGN_PARAGRAPH.CENTER


def _add_section(doc, sec, questions, section_no):
    title = SECTION_TITLES[sec["type"]]
    heading = doc.add_paragraph()
    run = heading.add_run(f"Section {section_no}: {title} ({sec['marks_each']} Mark(s) each)")
    _set_font(run, 12, True)

    choice_leave = sec.get("choice_leave", 0)
    if choice_leave:
        answer_n = sec["display_count"] - choice_leave
        note = doc.add_paragraph()
        note_run = note.add_run(f"(Answer any {answer_n} out of {sec['display_count']} questions)")
        note_run.italic = True
        _set_font(note_run, 10, False)

    for i, q in enumerate(questions, start=1):
        if sec["type"] == "mcq":
            table = doc.add_table(rows=1, cols=2)
            table.autofit = False
            table.columns[0].width = QUESTION_COL_WIDTH
            table.columns[1].width = MARK_COL_WIDTH
            q_cell, mark_cell = table.rows[0].cells
            q_cell.width = QUESTION_COL_WIDTH
            mark_cell.width = MARK_COL_WIDTH

            qp = q_cell.paragraphs[0]
            qp.paragraph_format.space_after = Pt(0)
            qrun = qp.add_run(f"{i}. {q['text']}")
            _set_font(qrun, 11, False)

            mp = mark_cell.paragraphs[0]
            mp.alignment = WD_ALIGN_PARAGRAPH.RIGHT
            mrun = mp.add_run("[    ]")
            _set_font(mrun, 11, False)

            after = doc.add_paragraph()
            after.paragraph_format.space_after = Pt(6)

            opts = q.get("_shuffled_options", q["options"])
            opt_text = "     ".join(
                f"({OPTION_LETTERS[i]}) {opt}" for i, opt in enumerate(opts)
            )
            op = doc.add_paragraph()
            op.paragraph_format.left_indent = Pt(18)
            orun = op.add_run(opt_text)
            _set_font(orun, 11, False)
        else:
            qp = doc.add_paragraph()
            qp.paragraph_format.space_after = Pt(6)
            qrun = qp.add_run(f"{i}. {q['text']}")
            _set_font(qrun, 11, False)
            if sec["type"] == "descriptive":
                doc.add_paragraph()  # a little writing space

    doc.add_paragraph()


def build_paper(fmt, meta, course_name, date_str, set_label, set_data, output_path):
    doc = Document()
    _set_a4_page(doc)
    style = doc.styles["Normal"]
    style.font.name = "Calibri"
    style.font.size = Pt(11)

    _header(doc, meta, course_name, fmt, date_str, set_label)

    for idx, sec in enumerate(fmt["sections"], start=1):
        _add_section(doc, sec, set_data[sec["type"]], idx)

    doc.save(output_path)
    return output_path


def build_answer_key(fmt, meta, course_name, date_str, set_label, set_data, output_path):
    doc = Document()
    _set_a4_page(doc)
    style = doc.styles["Normal"]
    style.font.name = "Calibri"
    style.font.size = Pt(11)

    p = doc.add_paragraph()
    p.alignment = WD_ALIGN_PARAGRAPH.CENTER
    r = p.add_run(f"ANSWER KEY — {course_name} — {fmt['name']} — Set {set_label} — {date_str}")
    _set_font(r, 13, True)
    doc.add_paragraph()

    for sec in fmt["sections"]:
        title = SECTION_TITLES[sec["type"]]
        heading = doc.add_paragraph()
        run = heading.add_run(title)
        _set_font(run, 12, True)

        for i, q in enumerate(set_data[sec["type"]], start=1):
            line = doc.add_paragraph()
            if sec["type"] == "mcq":
                text = f"{i}. {q['id']} — Correct: {q['correct_option']}"
            elif sec["type"] == "fill_blank":
                text = f"{i}. {q['id']} — Answer: {q['answer']}"
            else:
                text = f"{i}. {q['id']} — (Topic: {q['topic']}, Level: {q['level']}) — see subject notes"
            lr = line.add_run(text)
            _set_font(lr, 10.5, False)
        doc.add_paragraph()

    doc.save(output_path)
    return output_path


def safe_filename(*parts) -> str:
    name = "_".join(str(p) for p in parts)
    for ch in '<>:"/\\|?*':
        name = name.replace(ch, "")
    return name.strip().replace(" ", "_")
