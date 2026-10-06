"""CSV / Excel import & export for questions, courses and topics.

Import always creates new rows (any ID column in the file is ignored)
so freshly-assigned IDs never collide with what's already in the
JSON store. Export always writes CSV or XLSX based on the file
extension the user picked in the save dialog.
"""
import csv

from openpyxl import Workbook, load_workbook

MCQ_COLUMNS = ["id", "text", "option_a", "option_b", "option_c", "option_d", "correct_option", "topic", "level"]
FIB_COLUMNS = ["id", "text", "answer", "topic", "level"]
DES_COLUMNS = ["id", "text", "topic", "level"]
COURSE_COLUMNS = ["id", "name"]
TOPIC_COLUMNS = ["id", "name"]

QUESTION_COLUMNS = {"mcq": MCQ_COLUMNS, "fill_blank": FIB_COLUMNS, "descriptive": DES_COLUMNS}


def _question_to_row(qtype, q):
    if qtype == "mcq":
        opts = q.get("options", ["", "", "", ""])
        opts = (opts + ["", "", "", ""])[:4]
        return [q["id"], q["text"], *opts, q.get("correct_option", ""), q.get("topic", ""), q.get("level", "")]
    if qtype == "fill_blank":
        return [q["id"], q["text"], q.get("answer", ""), q.get("topic", ""), q.get("level", "")]
    return [q["id"], q["text"], q.get("topic", ""), q.get("level", "")]


def _row_to_question_fields(qtype, row_dict):
    """row_dict: column name -> value (id ignored, caller assigns a new one)."""
    if qtype == "mcq":
        options = [
            row_dict.get("option_a", "").strip(),
            row_dict.get("option_b", "").strip(),
            row_dict.get("option_c", "").strip(),
            row_dict.get("option_d", "").strip(),
        ]
        options = [o for o in options if o]
        return {
            "text": row_dict.get("text", "").strip(),
            "options": options,
            "correct_option": row_dict.get("correct_option", "").strip() or "NIL",
            "topic": row_dict.get("topic", "").strip(),
            "level": row_dict.get("level", "").strip() or "Medium",
        }
    if qtype == "fill_blank":
        return {
            "text": row_dict.get("text", "").strip(),
            "answer": row_dict.get("answer", "").strip() or "NIL",
            "topic": row_dict.get("topic", "").strip(),
            "level": row_dict.get("level", "").strip() or "Medium",
        }
    return {
        "text": row_dict.get("text", "").strip(),
        "topic": row_dict.get("topic", "").strip(),
        "level": row_dict.get("level", "").strip() or "Medium",
    }


def export_rows(path: str, columns: list, rows: list):
    if path.lower().endswith(".xlsx"):
        wb = Workbook()
        ws = wb.active
        ws.append(columns)
        for row in rows:
            ws.append(row)
        wb.save(path)
    else:
        with open(path, "w", newline="", encoding="utf-8-sig") as f:
            writer = csv.writer(f)
            writer.writerow(columns)
            writer.writerows(rows)


def _read_rows(path: str):
    """Yields dicts of column_name -> str value, header matched case-insensitively."""
    if path.lower().endswith(".xlsx"):
        wb = load_workbook(path, read_only=True, data_only=True)
        ws = wb.active
        rows_iter = ws.iter_rows(values_only=True)
        header = [str(h).strip().lower() if h is not None else "" for h in next(rows_iter)]
        for raw in rows_iter:
            if raw is None or all(v is None for v in raw):
                continue
            values = ["" if v is None else str(v) for v in raw]
            yield dict(zip(header, values))
    else:
        with open(path, "r", newline="", encoding="utf-8-sig") as f:
            reader = csv.reader(f)
            header = [h.strip().lower() for h in next(reader)]
            for raw in reader:
                if not raw or all(not v.strip() for v in raw):
                    continue
                yield dict(zip(header, raw))


def export_questions(qtype: str, questions: list, path: str):
    columns = QUESTION_COLUMNS[qtype]
    rows = [_question_to_row(qtype, q) for q in questions]
    export_rows(path, columns, rows)


def import_questions(qtype: str, path: str):
    """Returns a list of field-dicts (no id) ready for DataManager.add_question."""
    return [_row_to_question_fields(qtype, row) for row in _read_rows(path)]


def export_courses(courses: list, path: str):
    export_rows(path, COURSE_COLUMNS, [[c["id"], c["name"]] for c in courses])


def import_courses(path: str):
    return [row.get("name", "").strip() for row in _read_rows(path) if row.get("name", "").strip()]


def export_topics(topics: list, path: str):
    export_rows(path, TOPIC_COLUMNS, [[t["id"], t["name"]] for t in topics])


def import_topics(path: str):
    return [row.get("name", "").strip() for row in _read_rows(path) if row.get("name", "").strip()]
